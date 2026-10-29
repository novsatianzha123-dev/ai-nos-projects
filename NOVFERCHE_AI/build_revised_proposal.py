from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
ASSET=ROOT/'tmp'/'source_assets'
OUT=ROOT/'output'
(OUT/'docx').mkdir(parents=True,exist_ok=True)
(OUT/'pdf').mkdir(parents=True,exist_ok=True)

# Composite transparent images on white so they render consistently in Word/PDF.
for src,dst in [('p1_1_Image16.png','logo.png'),('p9_1_Image36.png','architecture.png'),('p11_1_Image42.png','flowchart.png')]:
    im=Image.open(ASSET/src).convert('RGBA')
    bg=Image.new('RGBA',im.size,'white'); bg.alpha_composite(im)
    bg.convert('RGB').save(OUT/'docx'/dst,quality=95)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(4); sec.left_margin=Cm(4); sec.bottom_margin=Cm(3); sec.right_margin=Cm(3)
sec.header_distance=Cm(1.5); sec.footer_distance=Cm(1.3)
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor(0,0,0)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [('Heading 1',12),('Heading 2',12)]:
    st=doc.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
    st.paragraph_format.space_before=Pt(4); st.paragraph_format.space_after=Pt(3); st.paragraph_format.keep_with_next=True
doc.styles['Title'].font.name='Arial'; doc.styles['Title'].font.size=Pt(14); doc.styles['Title'].font.bold=True; doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')

hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; hp.paragraph_format.space_after=Pt(0)
r=hp.add_run('PEKAN INSINYUR PETERNAKAN UGM 2026'); r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run('NOVFERCHE AI  |  Gagasan Inovasi Prototype'); r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

def p(text='',bold=False,align=None,size=None,after=4,before=0,italic=False,keep=False):
    x=doc.add_paragraph(); x.paragraph_format.space_before=Pt(before); x.paragraph_format.space_after=Pt(after)
    if align is not None: x.alignment=align
    if keep: x.paragraph_format.keep_with_next=True
    rr=x.add_run(text); rr.bold=bold; rr.italic=italic; rr.font.name='Arial'
    if size: rr.font.size=Pt(size)
    return x
def h(text,level=1): doc.add_paragraph(text,style=f'Heading {level}')
def nextpage(): doc.add_page_break()
def shade(cell,color):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(shd)
def celltext(cell,text,bold=False,size=9,align=None):
    cell.text=''; q=cell.paragraphs[0]; q.paragraph_format.space_after=Pt(0); q.paragraph_format.line_spacing=1.05
    if align is not None: q.alignment=align
    rr=q.add_run(str(text)); rr.font.name='Arial'; rr.font.size=Pt(size); rr.bold=bold
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def caption_table(number,title):
    p(f'Tabel {number}. {title}',bold=True,size=9.5,after=2,keep=True)
def table(headers,rows,widths=None,size=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers): celltext(t.rows[0].cells[i],x,True,size,WD_ALIGN_PARAGRAPH.CENTER); shade(t.rows[0].cells[i],'E5ECEF')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            celltext(cells[i],val,False,size)
            if ri%2: shade(cells[i],'F6F8F9')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    return t
def image(path,width_cm,caption,style='Gambar'):
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(1); q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Cm(width_cm))
    p(caption,size=8.5,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=5)

# Sampul
for _ in range(2): p('',after=8)
p('GAGASAN INOVASI PROTOTYPE',True,WD_ALIGN_PARAGRAPH.CENTER,14,after=4)
p('PEKAN INSINYUR PETERNAKAN UGM 2026',True,WD_ALIGN_PARAGRAPH.CENTER,11,after=18)
p('NOVFERCHE AI',True,WD_ALIGN_PARAGRAPH.CENTER,22,after=7)
p('Prototipe Pemantauan Kondisi Sapi Perah sebagai Dukungan\nKewaspadaan Awal Mastitis',True,WD_ALIGN_PARAGRAPH.CENTER,14,after=12)
p('Topik: Artificial Intelligence | IoT | Precision Livestock Farming',False,WD_ALIGN_PARAGRAPH.CENTER,10,after=18)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(10)
q.add_run().add_picture(str(OUT/'docx'/'logo.png'),width=Cm(3.0))
for line in ['Nama Ketua Tim: Novsa Ori Saputra','Nama Anggota: Feri Rahmansyah; Chelsea Ruth Christine Sipahutar','Nama Instansi: Telkom University','Program Studi: S1 Teknik Biomedis','Kabupaten/Kota: Bandung']:
    p(line,False,WD_ALIGN_PARAGRAPH.CENTER,10,after=2)
for _ in range(2): p('',after=4)
p('Yogyakarta\n2026',False,WD_ALIGN_PARAGRAPH.CENTER,10,after=0)

# Pengesahan
nextpage(); p('HALAMAN PENGESAHAN',True,WD_ALIGN_PARAGRAPH.CENTER,14,after=12)
p('Judul Prototype: NOVFERCHE AI - Prototipe Pemantauan Kondisi Sapi Perah sebagai Dukungan Kewaspadaan Awal Mastitis',after=5)
p('Topik: Artificial Intelligence, IoT, Precision Livestock Farming',after=3)
p('Nama Universitas: Telkom University',after=9)
h('Susunan Tim')
caption_table(1,'Susunan tim pengusul')
table(['No.','Nama','Peran'],[['1','Novsa Ori Saputra','Ketua'],['2','Feri Rahmansyah','Anggota'],['3','Chelsea Ruth Christine Sipahutar','Anggota']],[1,8,5],9)
p('Gagasan inovasi ini disusun untuk mengikuti Lomba Prototype Pekan Insinyur Peternakan UGM 2026.',before=5,after=16)
caption_table(2,'Ruang tanda tangan pengesahan')
table(['Ketua Tim','Mengetahui, Kaprodi'],[['\n\n\n(tanda tangan)\n\nNovsa Ori Saputra\nNIS/NIM: 101052430009','\n\n\n(tanda tangan)\n\nDr. Inung Wijayanto\nNIP/NIDN: 12860076']],[7,9],9)

# Orisinalitas
nextpage(); p('PERNYATAAN ORISINALITAS DAN STATUS PENGHARGAAN',True,WD_ALIGN_PARAGRAPH.CENTER,12,after=3)
p('KARYA LOMBA PROTOTYPE PEKAN INSINYUR PETERNAKAN UGM 2026',True,WD_ALIGN_PARAGRAPH.CENTER,10,after=12)
for line in ['Nama Ketua Tim: Novsa Ori Saputra','NIS/NIM: 101052430009','Nama Perguruan Tinggi: Telkom University','Program Studi: S1 Teknik Biomedis','Judul Prototype: NOVFERCHE AI - Prototipe Pemantauan Kondisi Sapi Perah sebagai Dukungan Kewaspadaan Awal Mastitis','Kategori: Sarjana/Profesi Insinyur']:
    p(line,after=2)
p('Dengan ini kami menyatakan bahwa karya/prototype yang diajukan dalam Lomba Prototype Pekan Insinyur Peternakan UGM 2026 bertema “Engineering Innovation for Sustainable Livestock Production” dikembangkan oleh tim kami dan:',before=6,after=3)
for line in [
'1. Belum pernah memperoleh Juara 1, 2, 3, atau penghargaan dalam lomba karya ilmiah, inovasi, prototype, atau kompetisi sejenis.',
'2. Merupakan karya asli yang dikembangkan tim, dapat dipertanggungjawabkan, dan tidak sedang diajukan sebagai karya pemenang pada kompetisi lain.',
'3. Menggunakan data, desain, gambar, dan sumber pihak lain secara jujur dengan mencantumkan rujukan.',
'4. Bersedia menerima konsekuensi sesuai ketentuan panitia apabila pernyataan ini terbukti tidak benar.'
]: p(line,after=3)
p('Demikian pernyataan ini dibuat dengan sebenar-benarnya untuk memenuhi persyaratan lomba.',before=3,after=8)
p('Bandung, 20 Oktober 2026',False,WD_ALIGN_PARAGRAPH.RIGHT,10,after=5)
caption_table(3,'Ruang tanda tangan pernyataan orisinalitas')
table(['Ketua Tim','Mengetahui, Kaprodi'],[['\n\n\n(tanda tangan)\n\nNovsa Ori Saputra\nNIS/NIM: 101052430009','\n\n\n(tanda tangan)\n\nDr. Inung Wijayanto\nNIP/NIDN: 12860076']],[7,9],9)

# Daftar Isi
nextpage(); p('DAFTAR ISI',True,WD_ALIGN_PARAGRAPH.CENTER,14,after=14)
for txt,n in [('PENDAHULUAN',5),('PERUMUSAN MASALAH',5),('MANFAAT INOVASI',6),('GAMBARAN INOVASI DAN TEKNOLOGI',6),('Nama dan Konsep Prototype',6),('Cara Kerja',6),('Komponen Utama',7),('Desain dan Alur Prototype',7),('Rencana Validasi',8),('Keunggulan Inovasi',9),('KESIMPULAN',9),('DAFTAR PUSTAKA',9),('LAMPIRAN TAMPILAN PROTOTYPE',10)]:
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.LEFT; q.paragraph_format.space_after=Pt(5); q.paragraph_format.line_spacing=1.0; q.paragraph_format.tab_stops.add_tab_stop(Cm(12))
    rr=q.add_run(f'{txt}\t{n}'); rr.font.name='Arial'; rr.font.size=Pt(10)

# Isi utama halaman 5
nextpage(); p('GAGASAN INOVASI PROTOTYPE',True,WD_ALIGN_PARAGRAPH.CENTER,14,after=2)
p('NOVFERCHE AI: Pemantauan Kondisi Sapi Perah untuk Mendukung Kewaspadaan Mastitis',True,WD_ALIGN_PARAGRAPH.CENTER,12,after=8)
h('PENDAHULUAN')
p('Di peternakan rakyat, pemantauan sapi perah masih banyak bergantung pada pemeriksaan langsung. Cara ini penting, tetapi tidak selalu kontinu atau tercatat sebagai data individual. Sensor wearable dapat merekam aktivitas, istirahat, perilaku makan, dan indikator fisiologis; ketepatannya bergantung pada jenis serta posisi sensor, kondisi lingkungan, dan validasi melalui pengamatan rujukan (Lee & Seo, 2021).')
p('Pembelajaran mesin telah dikaji untuk mengenali kejadian perawatan mastitis pada data peternakan (Post et al., 2020), tetapi hasil satu kumpulan data tidak otomatis berlaku di peternakan lain. NOVFERCHE AI karena itu diajukan sebagai prototipe pemantauan umum untuk mendukung kewaspadaan awal, bukan alat diagnosis. Kesesuaian klinis memerlukan data berlabel dan penilaian profesional.')
p('Perancangan fitur perilaku mengacu pada dataset akselerometer sapi Japanese Black yang memuat rekaman enam ekor sapi dan label 13 perilaku pada 25 Hz (Ito et al., 2022). Dataset tersebut bukan data sapi perah berlabel mastitis; evaluasi model prototipe tetap menggunakan 6.000 sampel simulasi. Karena itu, hasil internal belum membuktikan kinerja lapangan. Prototipe mengolah suhu tubuh, aktivitas, indeks gerakan, dan durasi istirahat menjadi kelas NORMAL atau PERLU PEMANTAUAN, lalu menampilkan indikator dalam tiga tingkat. Sensor fisik dan gateway belum terintegrasi.')
h('PERUMUSAN MASALAH')
p('Rumusan masalah: bagaimana merangkum data sensor individual sebagai sinyal yang mudah ditindaklanjuti; menghubungkan sensor/gateway, pemrosesan fitur, model, dan dashboard; serta mengevaluasi sistem dengan data lapangan berlabel dan pengamatan rujukan? Prototipe membantu memprioritaskan pemeriksaan, bukan mendiagnosis penyakit atau menggantikan tenaga kesehatan hewan.')

# Halaman 6
nextpage(); h('MANFAAT INOVASI')
p('Manfaat berikut merupakan harapan yang perlu diverifikasi melalui uji lapangan.',italic=True,size=9,after=3)
caption_table(4,'Manfaat inovasi yang diharapkan menurut penerima')
table(['Aspek','Manfaat yang diharapkan'],[
['Ternak','Perubahan aktivitas atau istirahat dapat terlihat lebih cepat melalui pencatatan individual, sehingga pemeriksaan dapat diprioritaskan.'],
['Peternak','Ringkasan data dan tingkat pemantauan yang mudah dibaca mendukung penentuan sapi yang perlu diperiksa serta melengkapi pencatatan rutin.'],
['Lingkungan','Dampak belum diukur. Penanganan lebih terarah berpotensi mengurangi penggunaan obat yang tidak perlu, tetapi hal ini harus diuji.'],
['Masyarakat/industri','Mendukung pengembangan praktik peternakan presisi yang memperhatikan kesehatan dan kesejahteraan ternak.'],
['Pengembangan teknologi','Menyediakan kerangka modular untuk pengujian sensor, gateway, dan data peternakan setempat.']],[3.1,12.9],8.2)
h('GAMBARAN INOVASI DAN TEKNOLOGI')
h('Nama dan Konsep Prototype',2)
p('NOVFERCHE AI adalah platform pemantauan kondisi sapi yang mengolah empat fitur sensor menjadi tingkat perhatian operasional. Perangkat lunak yang sudah tersedia mencakup model klasifikasi, dashboard Streamlit, dan API lokal FastAPI. Integrasi sensor fisik, firmware, dan gateway yang sesuai merupakan pekerjaan lanjutan.')
h('Cara Kerja',2)
caption_table(5,'Alur input, proses, teknologi, dan output prototipe')
table(['Input','Proses','Teknologi','Output'],[
['Suhu tubuh, aktivitas, indeks gerakan, durasi istirahat, ID sapi, dan waktu','Validasi rentang, penerimaan data, klasifikasi probabilitas kelas 1','Sensor/gateway (rencana), API FastAPI lokal, Random Forest, dashboard Streamlit','Indikator 0-100, kategori pemantauan, data terbaru, dan riwayat singkat']],[3.4,4.2,4.8,3.6],8)
p('API menerima fitur yang sudah diolah, bukan sinyal akselerometer mentah. Firmware/gateway perlu menghitung aktivitas, gerakan, dan durasi istirahat sebelum mengirim data. Dashboard juga dapat menerima input manual.',italic=True,size=8.5,after=0)

# Halaman 7
nextpage(); h('Komponen Utama')
caption_table(6,'Komponen dan fungsi dalam prototipe')
table(['No.','Komponen/teknologi','Fungsi'],[
['1','Sensor suhu dan akselerometer (rencana)','Mengukur suhu dan gerak ternak; jenis dan posisi pemasangan dipilih saat rekayasa lapangan.'],
['2','Firmware/gateway (rencana)','Mengolah bacaan menjadi fitur aktivitas, gerakan, dan istirahat serta meneruskannya ke API.'],
['3','API FastAPI lokal','Memvalidasi ID dan rentang data serta menyimpan pembacaan ke CSV lokal.'],
['4','Model Random Forest','Menghasilkan probabilitas kelas PERLU PEMANTAUAN dari empat fitur input.'],
['5','Dashboard Streamlit','Menampilkan data, indikator 0-100, kategori pemantauan, dan riwayat terbaru.']],[1,5,10],8.3)
h('Desain dan Alur Prototype')
p('Gambar 1 memperlihatkan aliran data dari masukan manual atau sensor uji sampai hasil pemantauan pada dashboard.',after=2)
image(OUT/'docx'/'architecture.png',15.5,'Gambar 1. Diagram arsitektur aliran data NOVFERCHE AI. Integrasi sensor dan gateway masih berupa rancangan; prototipe saat ini menerima fitur melalui input manual atau API lokal.')
p('Model mengubah probabilitas kelas 1 menjadi indikator 0-100. Dashboard menampilkan PEMANTAUAN RUTIN (<30), PERLU PEMANTAUAN (30 sampai <70), atau PEMANTAUAN LEBIH INTENSIF (>=70). Ambang ini merupakan aturan tampilan prototipe, bukan batas klinis atau diagnosis. Penerapannya untuk keputusan kesehatan harus ditinjau bersama dokter hewan dan prosedur peternakan.',after=0)

# Halaman 8
nextpage(); h('Alur Operasional dan Rencana Validasi')
p('Flowchart 1 menunjukkan pemeriksaan sumber, kelengkapan data, dan jalur keluaran kategori pada aplikasi.',after=3)
image(OUT/'docx'/'flowchart.png',5.3,'Flowchart 1. Alur operasional NOVFERCHE AI dari pemilihan sumber data sampai tampilan kategori pemantauan.')
h('Rencana Validasi')
p('Pengujian lanjutan perlu dilakukan bersama mitra peternakan dan tenaga kesehatan hewan. Ukuran keberhasilan dan prosedur ditetapkan sebelum pengambilan data.',after=3)
caption_table(7,'Tahapan validasi teknis dan pengguna')
table(['Tahap','Kegiatan dan ukuran keberhasilan'],[
['Integrasi perangkat','Pilih sensor dan lokasi pemasangan; periksa kalibrasi, kenyamanan, ketahanan, daya, dan kestabilan koneksi.'],
['Akuisisi dan pelabelan','Catat ID sapi, waktu, sensor, serta pengamatan rujukan oleh petugas terlatih; untuk sapi perah dapat dipertimbangkan California Mastitis Test sesuai protokol veteriner.'],
['Evaluasi model','Pisahkan data uji berdasarkan individu atau periode; laporkan confusion matrix, precision, recall, F1, dan sensitivitas kelas yang perlu perhatian.'],
['Uji penggunaan','Nilai keterbacaan dashboard, kemudahan alur, waktu respons, dan kesesuaian keluaran untuk keputusan pemantauan.'],
['Perbaikan','Tinjau kualitas data, ambang, kesalahan klasifikasi, dan batas penggunaan sebelum uji lebih luas.']],[3.1,12.9],8)

# Halaman 9
nextpage(); h('Evaluasi Internal dan Keunggulan Inovasi')
p('Uji internal pada 6.000 sampel simulasi (4.800 latih; 1.200 uji) menghasilkan akurasi 1,0000 dan confusion matrix [[853, 0], [0, 347]]. Label dibentuk dari rentang fitur sintetis yang mudah dipisahkan. Hasil ini bukan bukti performa pada sapi nyata; diperlukan validasi eksternal dengan data lapangan berlabel independen.')
h('Keunggulan Inovasi')
p('Nilai tambah NOVFERCHE AI adalah alur modular yang mengintegrasikan sensor, API, klasifikasi, dan dashboard untuk membantu prioritas pemeriksaan. Kebaruan terletak pada rancangan integrasi, bukan algoritma Random Forest. Keterjangkauan belum terukur dan baru dapat dinilai setelah komponen fisik dipilih.')
h('KESIMPULAN')
p('NOVFERCHE AI mengolah empat fitur sapi melalui API lokal dan model Random Forest, lalu menampilkan tiga tingkat pemantauan pada dashboard. Perangkat lunak telah berjalan, tetapi integrasi sensor fisik dan uji lapangan masih diperlukan. Karena model memakai data sintetis tanpa label mastitis, keluarannya adalah indikator umum, bukan diagnosis atau deteksi klinis. Pengembangan selanjutnya memerlukan data sapi perah berlabel, pengamatan rujukan, dan evaluasi independen.')
h('DAFTAR PUSTAKA')
for i,ref in enumerate([
'Lee, M., & Seo, S. (2021). Wearable Wireless Biosensor Technology for Monitoring Cattle: A Review. Animals, 11(10), 2779. https://doi.org/10.3390/ani11102779',
'Post, C., Rietz, C., Büscher, W., & Müller, U. (2020). Using Sensor Data to Detect Lameness and Mastitis Treatment Events in Dairy Cows: A Comparison of Classification Models. Sensors, 20(14), 3863. https://doi.org/10.3390/s20143863',
'Trendov, N. M., Varas, S., & Zeng, M. (2019). Digital technologies in agriculture and rural areas: Status report. Rome: Food and Agriculture Organization of the United Nations.',
'Ito, H., Takeda, K., Tokgoz, K. K., Minati, L., Fukawa, M., Chao, L., Bartels, J., Rachi, I., & A, S. (2022). Japanese Black Beef Cow Behavior Classification Dataset (Version 2.0.0) [Data set]. Zenodo. https://doi.org/10.5281/zenodo.5849025. Lisensi: CC BY-NC-ND 4.0.'
],1):
    q=doc.add_paragraph(); q.paragraph_format.left_indent=Cm(.55); q.paragraph_format.first_line_indent=Cm(-.55); q.paragraph_format.line_spacing=1.0; q.paragraph_format.space_after=Pt(1)
    rr=q.add_run(f'{i}. {ref}'); rr.font.name='Arial'; rr.font.size=Pt(9)

# Lampiran: beri caption untuk setiap screenshot dari prototipe.
nextpage(); p('LAMPIRAN TAMPILAN PROTOTYPE',True,WD_ALIGN_PARAGRAPH.CENTER,12,after=8)
image(ASSET/'p12_1_Image46.jpg',13.5,'Gambar 2. Tampilan awal dashboard NOVFERCHE AI sebelum klasifikasi dijalankan.')
image(ASSET/'p12_2_Image47.jpg',13.5,'Gambar 3. Contoh keluaran pemantauan rutin dengan indikator prototipe 25/100.')
image(ASSET/'p12_3_Image48.jpg',13.5,'Gambar 4. Contoh keluaran PERLU PEMANTAUAN dengan indikator prototipe 50/100.')

nextpage(); p('LAMPIRAN TAMPILAN PROTOTYPE (LANJUTAN)',True,WD_ALIGN_PARAGRAPH.CENTER,11,after=8)
image(ASSET/'p13_1_Image51.jpg',13.5,'Gambar 5. Contoh keluaran PEMANTAUAN LEBIH INTENSIF dengan indikator prototipe 100/100.')
image(ASSET/'p13_2_Image52.jpg',13.5,'Gambar 6. Dokumentasi antarmuka API lokal NOVFERCHE AI untuk menerima dan menampilkan data sensor.')

doc.core_properties.title='NOVFERCHE AI - Gagasan Inovasi Prototype'
doc.core_properties.subject='Pekan Insinyur Peternakan UGM 2026'
doc.core_properties.author='Tim NOVFERCHE AI'
doc.save(OUT/'docx'/'NovsaOriSaputra_TelkomUniversity_NOVFERCHEAI.docx')
print(OUT/'docx'/'NovsaOriSaputra_TelkomUniversity_NOVFERCHEAI.docx')
